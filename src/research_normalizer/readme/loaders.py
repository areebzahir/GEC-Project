"""Load a documentation file into normalized lines (plus tables for DOCX).

Each loader returns :class:`LoadedDocument` with a list of lines (1-based line numbers are the index
plus one) so every later extraction can cite a line range. Text is cleaned minimally here: Unicode
NFKC normalization and stripped trailing whitespace; the heavier quote-noise cleaning happens in the
segmenter where it has block context (DESIGN.md section 11).
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from ..config import PipelineConfig
from ..issues import Issue, IssueCode, Severity
from ..text_decoding import decode_bytes


@dataclass(slots=True)
class LoadedDocument:
    """Normalized lines of a documentation file, plus any extracted tables (DOCX)."""

    lines: list[str]
    encoding: str
    doc_format: str                         # markdown | text | pdf | docx
    tables: list[list[list[str]]] = field(default_factory=list)  # docx: list of row-lists
    issues: list[Issue] = field(default_factory=list)


def _normalize_line(line: str) -> str:
    # NFKC folds exotic look-alikes (e.g. full-width punctuation) to canonical forms, which makes
    # heading/key detection reliable across copy-pasted documents.
    return unicodedata.normalize("NFKC", line).replace("\r", "").rstrip()


def load_document(path: Path, extension: str, config: PipelineConfig) -> LoadedDocument:
    """Dispatch to the right loader by extension; never raises for a bad file."""
    if extension == ".pdf":
        return _load_pdf(path, config)
    if extension == ".docx":
        return _load_docx(path)
    return _load_text(path, extension, config)


def _load_text(path: Path, extension: str, config: PipelineConfig) -> LoadedDocument:
    decoded = decode_bytes(path.read_bytes(), config)
    lines = [_normalize_line(ln) for ln in decoded.text.split("\n")]
    fmt = "markdown" if extension in {".md", ".markdown"} else "text"
    issues: list[Issue] = []
    if decoded.rung in {"cp1252", "latin-1"}:
        issues.append(Issue.make(
            IssueCode.ENCODING_FALLBACK, Severity.INFO,
            f"Read {path.name} as {decoded.encoding} (not UTF-8).",
            technical_detail=f"decode rung={decoded.rung}",
            file=path.name,
        ))
    return LoadedDocument(lines=lines, encoding=decoded.encoding, doc_format=fmt, issues=issues)


def _load_pdf(path: Path, config: PipelineConfig) -> LoadedDocument:
    """Extract text per page with layout mode, which keeps column alignment (Ref [9])."""
    import pypdf

    issues: list[Issue] = []
    lines: list[str] = []
    try:
        reader = pypdf.PdfReader(str(path))
        pages = reader.pages[: config.max_pdf_pages]
        if len(reader.pages) > config.max_pdf_pages:
            issues.append(Issue.make(
                IssueCode.README_UNREADABLE, Severity.WARNING,
                f"{path.name} has more pages than the limit; only the first {config.max_pdf_pages} were read.",
                file=path.name,
            ))
        for number, page in enumerate(pages, start=1):
            text = page.extract_text(extraction_mode="layout") or ""
            lines.append(f"[page {number}]")
            lines.extend(_normalize_line(ln) for ln in text.split("\n"))
    except Exception as exc:  # noqa: BLE001 - pypdf raises various errors on bad PDFs
        issues.append(Issue.make(
            IssueCode.README_UNREADABLE, Severity.WARNING,
            f"Could not extract text from {path.name}.",
            technical_detail=f"{type(exc).__name__}: {str(exc).splitlines()[0][:120]}",
            file=path.name,
        ))
    return LoadedDocument(lines=lines, encoding="binary-pdf", doc_format="pdf", issues=issues)


def _load_docx(path: Path) -> LoadedDocument:
    """Read Word paragraphs (heading styles kept as headings) and tables as row lists."""
    import docx  # python-docx

    issues: list[Issue] = []
    lines: list[str] = []
    tables: list[list[list[str]]] = []
    try:
        document = docx.Document(str(path))
        for para in document.paragraphs:
            text = _normalize_line(para.text)
            # Mark heading-styled paragraphs so the segmenter treats them as headings even without
            # Markdown syntax (python-docx exposes the style name, e.g. "Heading 1").
            style = (para.style.name or "").lower() if para.style else ""
            if text and style.startswith("heading"):
                lines.append(f"## {text}")
            else:
                lines.append(text)
        for table in document.tables:
            rows = [[_normalize_line(cell.text) for cell in row.cells] for row in table.rows]
            tables.append(rows)
    except Exception as exc:  # noqa: BLE001
        issues.append(Issue.make(
            IssueCode.README_UNREADABLE, Severity.WARNING,
            f"Could not read {path.name}.",
            technical_detail=f"{type(exc).__name__}: {str(exc).splitlines()[0][:120]}",
            file=path.name,
        ))
    return LoadedDocument(lines=lines, encoding="binary-docx", doc_format="docx", tables=tables, issues=issues)
